"""Spreadsheet sources without OAuth.

  google_sheet_csv(http, sheet_id, tab)   a Google Sheet tab as CSV via the share
                                          link ("anyone with the link can view");
                                          `tab` is the tab name or the gid from the
                                          tab's URL
  onedrive_download(http, share_url)      a OneDrive / SharePoint "anyone with the
                                          link" share URL, resolved to the file bytes
  fetch(http, location)                   dispatch: Google share link, OneDrive
                                          share link, plain https URL, or local path
  read_table(data, filename, sheet)       rows as dicts from CSV or XLSX bytes with
                                          headers normalised (lower snake_case)

The template never stores your links: they live in .env (see config.py).
"""
from __future__ import annotations

import base64
import csv
import io
import re
from datetime import date, datetime
from pathlib import Path

import httpx

from app.integrations.events.common import ProviderError

GOOGLE_RE = re.compile(r"docs\.google\.com/spreadsheets/d/([A-Za-z0-9_-]+)")
ONEDRIVE_HOSTS = ("1drv.ms", "onedrive.live.com", "sharepoint.com", "my.sharepoint.com")


def norm_header(h) -> str:
    h = re.sub(r"\[merged\]", "", str(h or "")).strip().casefold()
    h = re.sub(r"[^a-z0-9]+", "_", h).strip("_")
    return h


def google_sheet_xlsx_url(sheet_id: str) -> str:
    """The whole workbook, every tab, as XLSX. Lets Aurora pick a tab by its headers."""
    return f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=xlsx"


def google_sheet_xlsx(http: httpx.Client, sheet_id: str) -> bytes:
    resp = _get(http, google_sheet_xlsx_url(sheet_id))
    if resp.content[:2] != b"PK":
        raise ProviderError("Google returned a sign-in page instead of the workbook: share the sheet as 'anyone with the link can view'")
    return resp.content


def google_sheet_csv_url(sheet_id: str, tab: str | None) -> str:
    base = f"https://docs.google.com/spreadsheets/d/{sheet_id}/gviz/tq?tqx=out:csv"
    if tab and tab.isdigit():
        return f"{base}&gid={tab}"
    if tab:
        return f"{base}&sheet={httpx.URL('', params={'s': tab}).params['s']}"
    return base


def _get(http: httpx.Client, url: str) -> httpx.Response:
    resp = http.get(url, follow_redirects=True)
    if resp.status_code >= 400:
        raise ProviderError(f"{url[:80]}: HTTP {resp.status_code}")
    return resp


def google_sheet_csv(http: httpx.Client, sheet_id: str, tab: str | None) -> bytes:
    url = google_sheet_csv_url(sheet_id, tab)
    resp = _get(http, url)
    ctype = resp.headers.get("content-type", "")
    if "text/html" in ctype or resp.content[:15].lower().startswith(b"<!doctype html"):
        raise ProviderError("Google returned a sign-in page: share the sheet as 'anyone with the link can view'")
    return resp.content


def onedrive_direct_url(share_url: str) -> str:
    """OneDrive's shares API accepts a base64url-encoded share link with a 'u!'
    prefix and serves the file at /root/content, no token needed for anonymous
    personal links."""
    token = base64.urlsafe_b64encode(share_url.encode()).decode().rstrip("=")
    return f"https://api.onedrive.com/v1.0/shares/u!{token}/root/content"


def sharepoint_download_url(share_url: str) -> str:
    """A OneDrive for Business / SharePoint 'anyone with the link' URL serves the
    file itself when `download=1` is appended."""
    u = httpx.URL(share_url)
    return str(u.copy_add_param("download", "1")) if "download" not in u.params else share_url


def onedrive_download(http: httpx.Client, share_url: str) -> bytes:
    host = httpx.URL(share_url).host or ""
    candidates = [sharepoint_download_url(share_url), onedrive_direct_url(share_url)] if "sharepoint" in host \
        else [onedrive_direct_url(share_url), sharepoint_download_url(share_url)]
    last = None
    for url in candidates:
        try:
            resp = _get(http, url)
        except ProviderError as e:
            last = e
            continue
        if resp.content[:2] == b"PK" or "spreadsheetml" in resp.headers.get("content-type", "") or "text/csv" in resp.headers.get("content-type", ""):
            return resp.content
        last = ProviderError("share link returned a web page, not the file: set the link to 'anyone with the link' and try again")
    raise last or ProviderError("could not download the share link")


def fetch(http: httpx.Client, location: str, tab: str | None = None) -> tuple[bytes, str]:
    """Returns (bytes, filename-ish) for a Google share link / OneDrive share link /
    https URL / local path."""
    m = GOOGLE_RE.search(location)
    if m:
        return google_sheet_csv(http, m.group(1), tab), "sheet.csv"
    if location.startswith("http"):
        host = httpx.URL(location).host or ""
        if any(h in host for h in ONEDRIVE_HOSTS):
            return onedrive_download(http, location), "workbook.xlsx"
        resp = _get(http, location)
        name = Path(httpx.URL(location).path).name or "download"
        ctype = resp.headers.get("content-type", "")
        if "spreadsheetml" in ctype and not name.endswith(".xlsx"):
            name += ".xlsx"
        return resp.content, name
    path = Path(location)
    if not path.exists():
        raise ProviderError(f"{location}: no such file")
    return path.read_bytes(), path.name


def read_table(data: bytes, filename: str, sheet: str | None = None, header_row: int | None = None) -> list[dict]:
    """Rows as dicts keyed by normalised header. Finds the header row automatically:
    the first row with at least three non-empty cells that are all text."""
    if filename.lower().endswith((".xlsx", ".xlsm")) or data[:2] == b"PK":
        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb[wb.sheetnames[0]]
        raw = [list(r) for r in ws.iter_rows(values_only=True)]
    else:
        text = data.decode("utf-8-sig", errors="replace")
        raw = [row for row in csv.reader(io.StringIO(text))]
    return rows_from_grid(raw, header_row)


def workbook_tabs(data: bytes) -> dict[str, list[dict]]:
    """Every worksheet of an XLSX as rows-by-tab-name."""
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = {}
    for name in wb.sheetnames:
        raw = [list(r) for r in wb[name].iter_rows(values_only=True)]
        out[name] = rows_from_grid(raw) if raw else []
    return out


def pick_tab(tabs: dict[str, list[dict]], required: tuple[str, ...], preferred: str | None = None) -> tuple[str, list[dict]]:
    """The tab named `preferred` if given, else the first tab whose headers contain
    every name in `required` (normalised)."""
    if preferred:
        for name, rows in tabs.items():
            if name.casefold().strip() == preferred.casefold().strip():
                return name, rows
        raise ProviderError(f"no tab named {preferred!r}; tabs are {list(tabs)}")
    for name, rows in tabs.items():
        if rows and all(r in rows[0] for r in required):
            return name, rows
    raise ProviderError(f"no tab has the columns {list(required)}; tabs are {list(tabs)}")


def rows_from_grid(raw: list[list], header_row: int | None = None) -> list[dict]:
    if header_row is None:
        header_row = 0
        for i, row in enumerate(raw[:50]):
            cells = [c for c in row if c not in (None, "")]
            if len(cells) >= 3 and all(isinstance(c, str) for c in cells) and not any(str(c).startswith("[merged]") for c in cells):
                header_row = i
                break
    headers = [norm_header(h) for h in raw[header_row]] if raw else []
    out = []
    for row in raw[header_row + 1:]:
        if not any(c not in (None, "") for c in row):
            continue
        d = {}
        for h, v in zip(headers, row):
            if h:
                d[h] = v
        out.append(d)
    return out


# ---------- cell helpers shared by the sheet importers ----------

def cell_str(v) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def cell_date(v) -> date | None:
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip().strip('"')
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%dT%H:%M:%S.%fZ", "%Y-%m-%dT%H:%M:%SZ",
                "%b %d, %Y", "%B %d, %Y", "%b %d, %Y %I:%M %p", "%d-%b-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    m = re.match(r"(\d{4}-\d{2}-\d{2})", s)
    if m:
        return date.fromisoformat(m.group(1))
    m = re.match(r"([A-Za-z]{3}) (\d{1,2}), (\d{4})", s)
    if m:
        try:
            return datetime.strptime(f"{m.group(1)} {m.group(2)} {m.group(3)}", "%b %d %Y").date()
        except ValueError:
            pass
    return None


def cell_num(v):
    from decimal import Decimal, InvalidOperation

    if v in (None, ""):
        return None
    if isinstance(v, (int, float)):
        return Decimal(repr(v)) if isinstance(v, float) else Decimal(v)
    s = re.sub(r"[,$%\s]", "", str(v))
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def cell_bool(v) -> bool:
    return str(v).strip().casefold() in ("true", "yes", "1", "y")
