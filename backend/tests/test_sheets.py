"""Spreadsheet sources: reader, OMMU market rows, competitor deals, promotions
workbook, sync orchestration. Network is httpx.MockTransport only."""
import io
import json
from datetime import date, datetime
from decimal import Decimal
from types import SimpleNamespace

import httpx
import openpyxl
import pytest
from sqlalchemy import select

from app.analytics.discounts import promotion_feed
from app.analytics.market import competitor_pressure, market_context
from app.analytics.money import pct_change
from app.analytics.periods import Period
from app.analytics.promotions import promotion_results
from app.importers.market import deals_to_events, import_market_rows, market_competition_events, offer_severity
from app.importers.promotions_sheet import active_days, import_promotion_rows, infer_type_value, parse_weekdays
from app.integrations.events.common import ProviderError
from app.integrations.events.sync import upsert_events
from app.integrations.sheets import google_sheet_csv_url, norm_header, onedrive_direct_url, read_table, rows_from_grid
from app.integrations.sheets_sync import sheets_sync
from app.models import DiscountDaily, ExternalEvent, MarketWeekly, PromoDayPerformance, Promotion, Store

OMMU_CSV = b"""row_id,week_ending,report_date,source_url,mmtc_name_canonical,is_p13_fl,dispensing_locations,medical_marijuana_mg_thc,low_thc_cannabis_mg_cbd,marijuana_smoking_oz,share_thc_pct,share_flower_pct,share_locations_pct,statewide_qualified_patients,is_totals_row
2026-09-03__trulieve,2026-09-03,9/4/2026,https://x/0903.pdf,Trulieve,FALSE,165,1.20E+08,0,50000,26.0,33.0,22.0,941000,FALSE
2026-09-03__p13,2026-09-03,9/4/2026,https://x/0903.pdf,Demo Operator Florida LLC,TRUE,34,7000000,0,3300,2.2,2.9,4.5,941000,FALSE
2026-09-03__sunburn,2026-09-03,9/4/2026,https://x/0903.pdf,Sunburn,FALSE,20,6500000,0,3100,2.1,2.4,2.7,941000,FALSE
2026-09-10__trulieve,2026-09-10,9/11/2026,https://x/0910.pdf,Trulieve,FALSE,165,1.10E+08,0,48000,25.0,32.0,22.0,941271,FALSE
2026-09-10__p13,2026-09-10,9/11/2026,https://x/0910.pdf,Demo Operator Florida LLC,TRUE,34,7350000,0,3400,2.29,3.19,4.5,941271,FALSE
2026-09-10__sunburn,2026-09-10,9/11/2026,https://x/0910.pdf,Sunburn,FALSE,24,7000000,0,3300,2.19,2.5,3.2,941271,FALSE
"""
DEALS_CSV = b"""row_id,deal_id,week_ending,operator_canonical,operator_display,observation_source,observed_at_utc,source_url_or_msg_id,subject,offer_type,offer_value,hook,audience,confidence
r1,d1,2026-09-10,Sunburn,Sunburn,flcd_site,2026-09-08T10:40:00.000Z,https://deals/1,SUNBURN 40% OFF FLOWER,bundle,40% off,urgency,general,high
r2,d2,2026-09-10,Mint Cannabis,Mint,flcd_site,2026-09-09T10:40:00.000Z,https://deals/2,MINT,value,(unparsed offer),,general,low
r3,d3,2026-09-10,Demo Operator Florida LLC,Demo Brand,flcd_site,2026-09-09T10:40:00.000Z,https://deals/3,OURS,value,30% off,,general,high
r4,d4,2026-09-10,Curaleaf Florida LLC,Curaleaf,gmail,2026-09-09T10:40:00.000Z,msg-9,Curaleaf BOGO,bogo,BOGO pre-rolls,urgency,members,medium
r5,d5,2026-09-10,GrowHealthy,GrowHealthy,flcd_site,,https://deals/5,no date,value,15% off,,general,high
r6,d6,2026-09-03,Sunburn,Sunburn,flcd_site,2026-08-30T10:40:00.000Z,https://deals/6,SUNBURN 20% OFF,value,20% off,,general,high
r7,d7,2026-09-10,,,flcd_site,2026-09-09T10:40:00.000Z,https://deals/7,FLCANNABIS DEALS banner,,,value,general,low
r8,d8,2026-09-10,Trulieve,Trulieve,flcd_site,2026-09-09T10:40:00.000Z,https://deals/8,,,,,general,low
,,2026-09-10,Sunburn,Sunburn,flcd_site,2026-09-08T10:40:00.000Z,https://deals/9,SUNBURN 10% OFF,percent_off,10% off,,general,high
,,2026-09-10,Sunburn,Sunburn,flcd_site,2026-09-10T10:40:00.000Z,https://deals/9,SUNBURN 10% OFF,percent_off,10% off,,general,high
"""


def mock(router):
    def handler(req):
        status, body, headers = router(req)
        return httpx.Response(status, content=body, headers=headers)
    return httpx.Client(transport=httpx.MockTransport(handler))


def settings(**over):
    base = dict(market_sheet_id="", market_sheet_tab="", market_self_operator="", deals_sheet_id="", deals_sheet_tab="",
                promotions_url="", promotions_sheet="", promotions_column_map="", market_centroid_lat=28.1, market_centroid_lon=-81.6)
    base.update(over)
    return SimpleNamespace(**base)


def promo_xlsx(rows, headers=("Promo", "Start Date", "End Date", "Discount", "Days", "Stores", "POS Discount Name", "Category", "Notes")):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Promos"
    ws.append(["Demo Brand FL promotions calendar"])  # a title row above the header
    ws.append([])
    ws.append(list(headers))
    for r in rows:
        ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------- reader ----------

def test_header_normalisation_and_grid_detection():
    assert norm_header(" POS Discount Name ") == "pos_discount_name"
    assert norm_header("[merged] Week of") == "week_of"
    grid = [["Title"], [], ["A", "B", "C"], [1, "", 3], [None, None, None], ["x", "y", "z"]]
    rows = rows_from_grid(grid)
    assert rows == [{"a": 1, "b": "", "c": 3}, {"a": "x", "b": "y", "c": "z"}]


def test_read_csv_and_xlsx():
    rows = read_table(OMMU_CSV, "m.csv")
    assert rows[0]["mmtc_name_canonical"] == "Trulieve" and len(rows) == 6
    data = promo_xlsx([("Try it Tuesday", date(2026, 9, 1), None, "60% OFF ALL 1G Distillate Vapes", "Tuesday", "", "DD - Auto - Try it Tuesday: 60% OFF ALL 1G Distillate Vapes!", "Vape Cart 1g", "")])
    rows = read_table(data, "promos.xlsx", "Promos")
    assert rows[0]["promo"] == "Try it Tuesday" and rows[0]["start_date"] == datetime(2026, 9, 1)
    assert "pos_discount_name" in rows[0]


def test_source_urls():
    assert google_sheet_csv_url("ID1", "1234") == "https://docs.google.com/spreadsheets/d/ID1/gviz/tq?tqx=out:csv&gid=1234"
    assert google_sheet_csv_url("ID1", "raw weekly").endswith("&sheet=raw weekly")
    u = onedrive_direct_url("https://1drv.ms/x/s!abc")
    assert u.startswith("https://api.onedrive.com/v1.0/shares/u!") and u.endswith("/root/content") and "=" not in u


def test_google_sign_in_page_is_an_error():
    with mock(lambda req: (200, b"<!DOCTYPE html><html>sign in</html>", {"content-type": "text/html"})) as http:
        from app.integrations.sheets import google_sheet_csv
        with pytest.raises(ProviderError, match="anyone with the link"):
            google_sheet_csv(http, "ID", None)


# ---------- market ----------

def test_market_rows_context_and_competition_events(session):
    rows = read_table(OMMU_CSV, "m.csv")
    r = import_market_rows(session, rows)
    assert (r.inserted, r.skipped, r.errors) == (6, 0, [])
    assert import_market_rows(session, rows).updated == 6
    me = session.execute(select(MarketWeekly).where(MarketWeekly.is_self == 1, MarketWeekly.week_ending == date(2026, 9, 10))).scalar_one()
    assert me.mg_thc == Decimal("7350000.00") and me.share_flower_pct == Decimal("3.1900")

    ctx = market_context(session, date(2026, 9, 15))
    assert ctx["current"]["week_ending"] == "2026-09-10"
    # market: 110M + 7.35M + 7M = 124.35M vs 120M + 7M + 6.5M = 133.5M -> -6.85%
    assert ctx["current"]["market_mg_thc"] == Decimal("124350000.00")
    assert ctx["vs_previous_week"]["market_mg_thc_pct"] == Decimal("-0.0685")
    assert ctx["vs_previous_week"]["self_mg_thc_pct"] == Decimal("0.0500")  # 7.35 / 7.0
    assert ctx["vs_previous_week"]["share_thc_bps"] == Decimal("9.0000")  # 2.29 - 2.20 = 0.09 pts = 9 bps
    assert ctx["vs_previous_week"]["share_flower_bps"] == Decimal("29.0000")
    assert ctx["vs_previous_week"]["market_dispensaries_delta"] == 4 and ctx["vs_previous_week"]["patients_delta"] == 271
    assert ctx["read"].startswith("Outpaced the market by 11.9 points")
    assert market_context(session, date(2026, 8, 1)) is None

    drafts = market_competition_events(session, (28.1, -81.6))
    assert len(drafts) == 1
    d = drafts[0]
    assert d.event_id == "ommu:2026-09-10:sunburn:+4" and d.severity == "major" and d.event_type == "competition"
    assert d.start_time == datetime(2026, 9, 4, 0, 0) and d.affected_radius_km == 1000.0
    assert "Sunburn added 4 dispensing locations (20 → 24)" in d.description


# ---------- deals ----------

def test_offer_severity():
    assert offer_severity("40% off", "bundle") == ("major", Decimal("40"))
    assert offer_severity("20% off", None) == ("moderate", Decimal("20"))
    assert offer_severity("15% off", None) == ("minor", Decimal("15"))
    assert offer_severity("BOGO pre-rolls", "bogo")[0] == "major"
    assert offer_severity("$25 eighths", None) == ("moderate", None)
    assert offer_severity("free stickers", "freebie") == ("minor", None)


def test_deals_to_events_and_pressure(session):
    session.add(Store(code="HS1", name="One", latitude=28.0, longitude=-82.0))
    session.commit()
    rows = read_table(DEALS_CSV, "d.csv")
    rep = deals_to_events(rows, (28.1, -81.6), "Demo Operator Florida LLC")
    # d2 has no parsed figure but a type and text: kept as a minor "value" deal. d7 has no
    # operator, d8 nothing at all. The two id-less Sunburn rows are one deal observed twice in a week.
    d = rep.to_dict()
    assert {k: d[k] for k in ("events", "skipped_self", "skipped_unparsed", "skipped_no_date", "skipped_no_operator", "duplicates")} == \
        {"events": 5, "skipped_self": 1, "skipped_unparsed": 1, "skipped_no_date": 1, "skipped_no_operator": 1, "duplicates": 1}
    by = {d.event_id: d for d in rep.drafts}
    assert by["deal:d2"].severity == "minor" and by["deal:d2"].description == "Mint Cannabis: value: MINT (value)"
    assert by["deal:d1"].severity == "major" and by["deal:d1"].confidence == Decimal("0.9")
    assert by["deal:d4"].severity == "major" and by["deal:d4"].confidence == Decimal("0.7")
    assert by["deal:d1"].start_time == datetime(2026, 9, 8) and by["deal:d1"].end_time == datetime(2026, 9, 15, 23, 59)
    assert upsert_events(session, rep.drafts, "deal-intel").inserted == 5

    week = Period("w", date(2026, 9, 7), date(2026, 9, 13))
    p = competitor_pressure(session, week)
    assert p["total_deals"] == 4 and p["previous_total_deals"] == 1 and p["vs_previous_pct"] == Decimal("3.0000")
    ops = {o["operator"]: o for o in p["operators"]}
    assert ops["Sunburn"]["deals"] == 2 and ops["Sunburn"]["previous_deals"] == 1 and ops["Sunburn"]["major"] == 1
    assert ops["Curaleaf Florida LLC"]["previous_deals"] == 0 and ops["Curaleaf Florida LLC"]["vs_previous_pct"] is None


# ---------- promotions ----------

def test_weekday_and_type_parsing():
    assert parse_weekdays("Tuesday") == "2" and parse_weekdays("tue, thu") == "2,4"
    assert parse_weekdays("Mon-Wed") == "1,2,3" and parse_weekdays("Fri - Mon") == "1,5,6,7"
    assert parse_weekdays("Daily") is None and parse_weekdays("") is None and parse_weekdays("weekends") == "6,7"
    assert infer_type_value(None, "60% OFF ALL 1G Distillate Vapes") == ("percent", Decimal("60"))
    assert infer_type_value(None, "$10 Off 3.5g & 7g DB Flower") == ("amount", Decimal("10"))
    assert infer_type_value("bogo", "Buy one get one pre-rolls") == ("bogo", Decimal("1"))
    assert infer_type_value("percent", 35) == ("percent", Decimal("35"))
    assert infer_type_value(None, 12.5) == ("percent", Decimal("12.5"))


def test_promotions_import_and_feed(session):
    session.add_all([Store(code="HS10001", name="FL - Demo - Pace"), Store(code="HS10002", name="FL - Demo - Tampa Kennedy")])
    session.commit()
    data = promo_xlsx([
        ("Try it Tuesday", date(2026, 9, 1), None, "60% OFF ALL 1G Distillate Vapes", "Tuesday", "", "DD - Auto - Try it Tuesday: 60% OFF ALL 1G Distillate Vapes!", "Vape Cart 1g", ""),
        ("Pace Grand Reopening", "9/12/2026", "9/13/2026", "$10 off", "", "Pace", "DD - Auto - Pace $10", "", "two days"),
        ("Bad row", None, None, "10%", "", "", "", "", ""),
        ("Nowhere deal", date(2026, 9, 1), date(2026, 9, 2), "5%", "", "Atlantis", "", "", ""),
        ("Pace Grand Reopening", "9/12/2026", "9/13/2026", "$10 off", "", "Tampa Kennedy", "DD - Auto - Pace $10", "", ""),
    ])
    rows = read_table(data, "promos.xlsx")
    r = import_promotion_rows(session, rows)
    assert (r.inserted, r.updated, r.skipped) == (3, 1, 1)
    assert any("no start date" in e for e in r.errors) and any("Atlantis" in e for e in r.errors)
    tt = session.execute(select(Promotion).where(Promotion.name == "Try it Tuesday")).scalar_one()
    assert (tt.discount_type, tt.discount_value, tt.weekdays, tt.store_codes) == ("percent", Decimal("60"), "2", None)
    assert tt.end_date == date(2027, 9, 1) and "open-ended" in tt.notes and tt.source == "sheet"
    assert tt.discount_names == "DD - Auto - Try it Tuesday: 60% OFF ALL 1G Distillate Vapes!"
    pg = session.execute(select(Promotion).where(Promotion.name == "Pace Grand Reopening")).scalar_one()
    assert sorted(pg.store_codes.split("|")) == ["HS10001", "HS10002"]  # two rows merged
    assert active_days(tt, date(2026, 9, 1), date(2026, 9, 30)) == [date(2026, 9, 1), date(2026, 9, 8), date(2026, 9, 15), date(2026, 9, 22), date(2026, 9, 29)]

    # feed join: two Tuesdays of Headset discount rows at Pace, plus a Wednesday row that must not count
    pace = session.execute(select(Store).where(Store.code == "HS10001")).scalar_one()
    for d, disc, rev, t in ((date(2026, 9, 8), "3000", "2000", 60), (date(2026, 9, 15), "3888", "2482.04", 70), (date(2026, 9, 16), "999", "1", 1)):
        session.add(DiscountDaily(store_id=pace.id, sale_date=d, discount_name="DD - Auto - Try it Tuesday: 60% OFF ALL 1G Distillate Vapes!",
                                  discount_total=Decimal(disc), revenue=Decimal(rev), units=100, transaction_count=t))
    session.commit()
    feed = promotion_feed(session, tt, "HS10001")
    assert feed["window"]["discount_total"] == Decimal("6888.00") and feed["window"]["days_with_data"] == 2
    assert feed["window"]["discount_per_day"] == Decimal("3444.00") and feed["window"]["tickets_per_day"] == Decimal("65.0000")
    assert feed["window"]["discount_depth"] == Decimal("0.6058")  # 6888 / (4482.04 + 6888)
    assert feed["baseline"] is None  # nothing before Sep 1
    assert promotion_feed(session, tt, "HS10002")["window"]["days_with_data"] == 0
    assert promotion_feed(session, pg, "HS10001")["applies_to_store"] is True
    results = {p["promotion"]: p for p in promotion_results(session, "HS10001")}
    assert results["Try it Tuesday"]["weekdays"] == "2" and results["Try it Tuesday"]["feed"]["window"]["discount_total"] == Decimal("6888.00")
    assert results["Try it Tuesday"]["verdict"] == "no_baseline"  # no line-level data: the POS path stays honest


def test_column_map_override(session):
    data = promo_xlsx([("X", date(2026, 1, 1), date(2026, 1, 2), "10%", "", "", "", "", "")], headers=("Deal Title", "Live From", "Live To", "Pct", "", "", "", "", ""))
    rows = read_table(data, "p.xlsx")
    r = import_promotion_rows(session, rows, json.dumps({"name": "Deal Title", "start": "Live From", "end": "Live To", "value": "Pct"}))
    assert r.inserted == 1
    p = session.execute(select(Promotion)).scalar_one()
    assert p.name == "X" and p.end_date == date(2026, 1, 2) and p.discount_type == "percent"


# ---------- sync ----------

def test_sheets_sync_end_to_end(session):
    session.add(Store(code="HS10001", name="FL - Demo - Pace", latitude=30.6, longitude=-87.16))
    session.commit()
    xlsx = promo_xlsx([("Try it Tuesday", date(2026, 9, 1), None, "60% off vapes", "Tue", "", "DD - Auto - Try it Tuesday", "", "")])

    def router(req):
        url = str(req.url)
        if "docs.google.com" in url and "MARKET" in url:
            assert "gid=77" in url
            return 200, OMMU_CSV, {"content-type": "text/csv"}
        if "docs.google.com" in url and "DEALS" in url:
            assert "sheet=deals" in url
            return 200, DEALS_CSV, {"content-type": "text/csv"}
        if "api.onedrive.com" in url:
            return 200, xlsx, {"content-type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}
        return 404, b"", {}

    cfg = settings(market_sheet_id="MARKET", market_sheet_tab="77", deals_sheet_id="DEALS", deals_sheet_tab="deals_library",
                   market_self_operator="Demo Operator Florida LLC",
                   promotions_url="https://demo-my.sharepoint.com/:x:/g/personal/x/abc?e=1", promotions_sheet="Promos")
    with mock(router) as http:
        rep = sheets_sync(session, http, cfg)
    assert rep.ran == ["market", "deals", "promotions"] and rep.warnings == [] and rep.skipped == {}
    assert rep.details["market"] == {"tab": "77", "rows": 6, "competition_events": 1}
    assert rep.details["deals"]["events"] == 5
    assert rep.details["promotions"]["rows"] == 1 and "pos_discount_name" in rep.details["promotions"]["headers"]
    kinds = {}
    for ev in session.execute(select(ExternalEvent)).scalars():
        kinds[ev.source] = kinds.get(ev.source, 0) + 1
    assert kinds == {"ommu": 1, "deal-intel": 5}
    assert session.execute(select(Promotion)).scalar_one().weekdays == "2"


def _workbook(tabs: dict[str, list[list]]) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, grid in tabs.items():
        ws = wb.create_sheet(name)
        for row in grid:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_sync_finds_tabs_by_headers_when_no_tab_is_configured(session):
    import csv as _csv
    market_grid = [r for r in _csv.reader(io.StringIO(OMMU_CSV.decode()))]
    deals_grid = [r for r in _csv.reader(io.StringIO(DEALS_CSV.decode()))]
    market_xlsx = _workbook({"Dashboard": [["Last refresh"], ["x", "y"]], "raw_weekly": market_grid, "Trends": [["week", "a", "b"]]})
    deals_xlsx = _workbook({"Briefing": [["hello"]], "deals_library": deals_grid})
    promo_xlsx_bytes = _workbook({"Cover": [["Demo Brand promos"]], "FL Deals": [["Promo", "Start Date", "End Date", "Discount", "Days"], ["Try it Tuesday", date(2026, 9, 1), None, "60% off", "Tue"]]})
    seen = []

    def router(req):
        url = str(req.url)
        seen.append(url)
        if "MARKET/export?format=xlsx" in url:
            return 200, market_xlsx, {"content-type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}
        if "DEALS/export?format=xlsx" in url:
            return 200, deals_xlsx, {"content-type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}
        if "sharepoint.com" in url and "download=1" in url:
            return 200, promo_xlsx_bytes, {"content-type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}
        return 404, b"", {}

    cfg = settings(market_sheet_id="MARKET", deals_sheet_id="DEALS", market_self_operator="Demo Operator Florida LLC",
                   promotions_url="https://netorg-my.sharepoint.com/:x:/g/personal/someone/IQBabc?e=nsvmdG")
    with mock(router) as http:
        rep = sheets_sync(session, http, cfg)
    assert rep.warnings == [] and rep.ran == ["market", "deals", "promotions"]
    assert rep.details["market"]["tab"] == "raw_weekly" and rep.details["market"]["rows"] == 6
    assert rep.details["deals"]["tab"] == "deals_library" and rep.details["deals"]["events"] == 5
    assert rep.details["promotions"]["tab"] == "FL Deals" and rep.details["promotions"]["rows"] == 1
    assert any("sharepoint.com" in u and "download=1" in u and "e=nsvmdG" in u for u in seen)
    assert session.execute(select(Promotion)).scalar_one().weekdays == "2"


def test_sync_reports_missing_tab_columns(session):
    wb = _workbook({"Only": [["a", "b", "c"], [1, 2, 3]]})
    with mock(lambda req: (200, wb, {"content-type": "application/octet-stream"})) as http:
        rep = sheets_sync(session, http, settings(market_sheet_id="M"), which={"market"})
    assert rep.ran == [] and "no tab has the columns" in rep.warnings[0] and "['Only']" in rep.warnings[0]


def test_sheets_sync_skips_unconfigured_and_reports_errors(session):
    with mock(lambda req: (403, b"forbidden", {})) as http:
        rep = sheets_sync(session, http, settings(deals_sheet_id="X"))
    assert set(rep.skipped) == {"market", "promotions"} and rep.ran == [] and "deals: " in rep.warnings[0] and "403" in rep.warnings[0]


def test_promotion_calendar_rows_are_single_days_and_long_names(session):
    """A promo calendar: one row per day, no end column, the day's deals packed
    into one cell. Each deal becomes its own one-day promotion with its own
    figure; the same deal on touching days merges into one window; the same deal
    weeks later is a separate window (unique per name+start)."""
    packed = ("55% Off All Edibles, Flower 1/8 oz, and 0.5ml Vapes; 50% Off All Pre-Rolls, 1ml Vapes, Tinctures, & Tablets; "
              "Door Buster: First 20 customers Get a Free Mac 1 3.5g Flower Jar with a $50+ Purchase")
    long_single = "Pull Tab Activation: Get 1 Pull Tab with a purchase $75 after discount to use on your next purchase, Get 2 pull tabs with a purchase of $125 after discount to use on next purchase and 1 to redeem immediately"
    data = promo_xlsx([
        (packed, date(2026, 1, 1), None, "", "", "", "", "", "New Years Day"),
        (long_single, date(2026, 1, 1), None, "", "", "", "", "", ""),
        ("Manager's Special", date(2026, 1, 5), None, "60% off", "", "", "", "", ""),
        ("Manager's Special", date(2026, 1, 6), None, "60% off", "", "", "", "", ""),
        ("Manager's Special", date(2026, 2, 3), None, "60% off", "", "", "", "", ""),
    ])
    r = import_promotion_rows(session, read_table(data, "cal.xlsx"))
    assert (r.inserted, r.updated, r.skipped, r.errors) == (6, 1, 0, [])
    day1 = session.execute(select(Promotion).where(Promotion.start_date == date(2026, 1, 1)).order_by(Promotion.id)).scalars().all()
    assert [(p.discount_type, p.discount_value) for p in day1] == [("percent", Decimal("55")), ("percent", Decimal("50")), ("amount", Decimal("50")), ("amount", Decimal("75"))]
    assert day1[0].name.startswith("55% Off") and day1[0].end_date == date(2026, 1, 1) and day1[0].notes == "New Years Day"
    assert len(day1[3].name) > 128
    ms = session.execute(select(Promotion).where(Promotion.name == "Manager's Special").order_by(Promotion.start_date)).scalars().all()
    assert [(p.start_date, p.end_date, p.discount_value) for p in ms] == [(date(2026, 1, 5), date(2026, 1, 6), Decimal("60")), (date(2026, 2, 3), date(2026, 2, 3), Decimal("60"))]
    # re-import is idempotent
    r2 = import_promotion_rows(session, read_table(data, "cal.xlsx"))
    assert (r2.inserted, r2.updated, r2.removed) == (0, 7, 0)
    assert len(session.execute(select(Promotion)).scalars().all()) == 6
    # a deal dropped from the sheet is dropped from Aurora; a manual one survives
    session.add(Promotion(name="Manual deal", start_date=date(2026, 3, 1), end_date=date(2026, 3, 1), discount_type="amount", discount_value=Decimal("0"), source="manual"))
    session.commit()
    data2 = promo_xlsx([("Manager's Special", date(2026, 2, 3), None, "60% off", "", "", "", "", "")])
    r3 = import_promotion_rows(session, read_table(data2, "cal.xlsx"))
    assert (r3.inserted, r3.updated, r3.removed) == (0, 1, 5)
    assert sorted(p.name for p in session.execute(select(Promotion)).scalars()) == ["Manager's Special", "Manual deal"]


def test_init_db_migrates_old_promotions_schema(tmp_path):
    """A database created by the earlier model (name VARCHAR(128) UNIQUE) must
    open, gain the (name, start_date) index and accept a long name on SQLite."""
    from sqlalchemy import create_engine, text
    from app.db import init_db
    eng = create_engine(f"sqlite:///{tmp_path/'old.db'}")
    with eng.begin() as c:
        c.execute(text("CREATE TABLE promotions (id INTEGER PRIMARY KEY, name VARCHAR(128) UNIQUE, start_date DATE, end_date DATE, "
                       "discount_type VARCHAR(16), discount_value NUMERIC(12,4), eligible_skus TEXT, eligible_category VARCHAR(64))"))
    applied = init_db(eng)
    assert any("ADD COLUMN weekdays" in a for a in applied)
    with eng.connect() as c:
        idx = [r[1] for r in c.execute(text("PRAGMA index_list(promotions)"))]
        assert "uq_promotions_name_start" in idx


def test_deals_recover_operator_from_text():
    from app.importers.market import operator_matcher
    m = operator_matcher(["Trulieve, Inc.", "Curaleaf Florida LLC", "Green Thumb Industries", "AltMed Florida"])
    assert m("VALID 9.8 curaleaf DISPENSARY Mention FL CANNABIS DEALS") == "Curaleaf Florida LLC"
    assert m("müv SEPTEMBER 8 - 10 40% OFF FLOWER") == "AltMed Florida"          # accent-insensitive brand alias
    assert m("RISE Dispensaries weekend deals") == "RISE Dispensaries"
    assert m("Jungle Boys drop") == "Jungle Boys" and m("FLCANNABIS DEALS.org FREE STICKERS!") is None
    assert m("Green Thumb 20% off") == "RISE Dispensaries"                        # multi-word alias beats the generic first word
    rows = read_table(
        b"deal_id,operator,observed_at_utc,full_deal_text_ocr,offer_type,offer_value,confidence\n"
        b"x1,,2026-09-09T10:40:00.000Z,VALID 9.8 curaleaf DISPENSARY 40% OFF STOREWIDE,percent_off,40,high\n"
        b"x2,,2026-09-09T10:40:00.000Z,FLCANNABIS DEALS banner,,,low\n"
        b",,,,,,\n", "d.csv")
    rep = deals_to_events(rows, (28.1, -81.6), None, known_operators=["Curaleaf Florida LLC"])
    d = rep.to_dict()
    assert (d["events"], d["recovered_operator"], d["skipped_no_operator"], d["skipped_blank"]) == (1, 1, 1, 0)  # read_table drops empty rows
    assert deals_to_events([{"operator": None, "observed_at_utc": None, "full_deal_text_ocr": ""}], (28.1, -81.6)).skipped_blank == 1
    assert rep.drafts[0].metadata["operator"] == "Curaleaf Florida LLC" and rep.drafts[0].severity == "major"
    assert d["unattributed_samples"] == ["2026-09-09 · FLCANNABIS DEALS banner"]


def test_promotion_year_fixed_from_weekday_column(session):
    from app.importers.promotions_sheet import fix_year_by_weekday
    assert fix_year_by_weekday(date(2025, 1, 1), "Thursday") == date(2026, 1, 1)   # Jan 1 2026 is a Thursday
    assert fix_year_by_weekday(date(2026, 1, 1), "Thursday") == date(2026, 1, 1)
    assert fix_year_by_weekday(date(2026, 1, 1), None) == date(2026, 1, 1)
    assert fix_year_by_weekday(date(2026, 1, 1), "Monday") == date(2026, 1, 1)      # neither neighbour fits: leave it
    rows = [
        {"january_daily_promos": "Thursday", "start_date": "2025-01-01", "promo": "55% Off Edibles", "notes": "NYD"},
        {"january_daily_promos": "Friday ", "start_date": "2025-01-02", "promo": "42% Off Storewide", "notes": None},
    ]
    r = import_promotion_rows(session, rows)
    assert r.inserted == 2 and any("date year corrected" in e for e in r.errors)
    assert sorted(p.start_date for p in session.execute(select(Promotion)).scalars()) == [date(2026, 1, 1), date(2026, 1, 2)]


def test_sync_reads_every_promo_performance_tab(session):
    grid = lambda rows: [["Start Date", "Promo", "Notes"], *rows]
    wb = _workbook({
        "Promo Performance 2025": grid([[date(2025, 7, 1), "50% Flower", ""]]),
        "Promo Performance 2026": grid([[date(2026, 1, 1), "55% Off Edibles; 42% Off Storewide", "NYD"]]),
        "Filtered Promo Performance": grid([[date(2026, 1, 1), "should not import twice", ""]]),
        "Promo Suggestions - March": grid([[date(2026, 3, 1), "not a performance tab", ""]]),
    })

    def router(req):
        return 200, wb, {"content-type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}

    cfg = settings(promotions_url="https://netorg-my.sharepoint.com/:x:/g/personal/x/abc?e=1")
    with mock(router) as http:
        rep = sheets_sync(session, http, cfg, which=("promotions",))
    assert rep.details["promotions"]["tab"] == "Promo Performance 2025, Promo Performance 2026"
    names = sorted(p.name for p in session.execute(select(Promotion)).scalars())
    assert names == ["42% Off Storewide", "50% Flower", "55% Off Edibles"]


def test_workbook_day_totals_are_gated_and_feed_the_statewide_block(session):
    from app.importers.promotions_sheet import import_promo_day_performance
    from app.analytics.promotions import promotion_results
    rows = [
        {"january_daily_promos": "Thursday", "start_date": "2025-01-01", "promo": "55% Off Edibles; 42% Off Storewide", "notes": "NYD",
         "net_sales": 57736.53, "gross_sales": 125343, "discount_amount": 67606.47, "promo_efficiency_roi": 0.854, "discount_rate": 0.539,
         "sales_per_hour": 9622.755, "4_week_average_sales": 79674, "forcasted_sales": 93550, "transaction_count": None},
        {"january_daily_promos": "Friday", "start_date": "2025-01-02", "promo": "Manager's Special", "notes": None,
         "net_sales": 101006.3, "gross_sales": 190179, "discount_amount": 89172.7, "promo_efficiency_roi": 1.1327, "4_week_average_sales": 103165},
        {"january_daily_promos": None, "start_date": None, "promo": "no date row"},
    ]
    # off by default: the sync leaves the table empty
    grid = [["January Daily Promos", "Start Date", "Promo", "Notes", "Net Sales", "Gross Sales", "Discount Amount", "Promo Efficiency ROI", "4 Week Average Sales"],
            ["Thursday", date(2025, 1, 1), "55% Off Edibles; 42% Off Storewide", "NYD", 57736.53, 125343, 67606.47, 0.854, 79674]]
    wb = _workbook({"Promo Performance 2026": grid})
    with mock(lambda req: (200, wb, {"content-type": "application/octet-stream"})) as http:
        rep = sheets_sync(session, http, settings(promotions_url="https://netorg-my.sharepoint.com/:x:/g/personal/x/abc?e=1"), which=("promotions",))
    assert rep.details["promotions"]["day_totals"].startswith("not imported")
    assert session.execute(select(PromoDayPerformance)).scalars().all() == []
    # switched on: one row per corrected date, and the promotion cards gain a statewide block
    r = import_promo_day_performance(session, rows, "Promo Performance 2026")
    assert (r.inserted, r.updated, r.skipped) == (2, 0, 1)
    perf = {p.day: p for p in session.execute(select(PromoDayPerformance)).scalars()}
    assert set(perf) == {date(2026, 1, 1), date(2026, 1, 2)} and perf[date(2026, 1, 1)].net_sales == Decimal("57736.53")
    assert perf[date(2026, 1, 1)].four_week_avg_sales == Decimal("79674.00") and perf[date(2026, 1, 2)].promo_roi == Decimal("1.1327")
    r2 = import_promo_day_performance(session, rows, "Promo Performance 2026")
    assert (r2.inserted, r2.updated) == (0, 2)
    import_promotion_rows(session, rows)   # the calendar rows themselves
    by = {p["promotion"]: p for p in promotion_results(session)}
    sw = by["42% Off Storewide"]["statewide"]
    assert sw["net_sales"] == Decimal("57736.53") and sw["discount_rate"] == Decimal("0.5394") and sw["days_with_data"] == 1
    assert sw["vs_four_week_pct"] == pct_change(Decimal("57736.53"), Decimal("79674"))
    assert by["Manager's Special"]["statewide"]["promo_roi"] == Decimal("1.1327")


def test_discount_codes_match_calendar_deals_automatically(session):
    from app.analytics.discounts import match_discount_names
    codes = ["DD - Auto - Try it Tuesday: 60% OFF ALL 1G Distillate Vapes!", "DD - Auto - 35% OFF ALL Demo Brand Flower!",
             "DD - Auto - 40% OFF ALL Demo Brand Derivative Products!", "CG - Auto - 40% Veterans & First Responders",
             "CG - Auto - First Time Patient", "DD - Auto - Manager's Special Menu: 60% Off"]
    assert match_discount_names("35% Off Demo Brand Flower - all sizes & all tiers", codes) == ["DD - Auto - 35% OFF ALL Demo Brand Flower!"]
    assert match_discount_names("40% Off Demo Brand Derivative Products Storewide", codes) == ["DD - Auto - 40% OFF ALL Demo Brand Derivative Products!"]
    assert match_discount_names("Try it Tuesday: 60% OFF ALL 1G Distillate Vapes", codes) == ["DD - Auto - Try it Tuesday: 60% OFF ALL 1G Distillate Vapes!"]
    assert match_discount_names("Manager's Special Menu: 60% Off Select Items", codes) == ["DD - Auto - Manager's Special Menu: 60% Off"]
    assert match_discount_names("Door Buster: free jar with $50 purchase", codes) == []      # no code carries that figure
    assert match_discount_names("40% Off Storewide", codes) == []                           # veterans code shares the figure, not the product

    # end to end: a promotion with no listed codes picks them up from the feed on its days
    session.add(Store(code="HS10001", name="FL - Demo - Pace", state="FL"))
    session.commit()
    pace = session.execute(select(Store)).scalar_one()
    session.add(DiscountDaily(store_id=pace.id, sale_date=date(2026, 9, 15), discount_name="DD - Auto - 35% OFF ALL Demo Brand Flower!",
                              discount_total=Decimal("715.75"), revenue=Decimal("1277.40"), units=48, transaction_count=33))
    session.add(DiscountDaily(store_id=pace.id, sale_date=date(2026, 9, 15), discount_name="CG - Auto - First Time Patient",
                              discount_total=Decimal("337"), revenue=Decimal("337"), units=25, transaction_count=8))
    import_promotion_rows(session, [{"promo": "35% Off Demo Brand Flower - all sizes & all tiers", "start_date": "2026-09-15"}])
    promo = session.execute(select(Promotion)).scalar_one()
    feed = promotion_feed(session, promo, "HS10001")
    assert feed["matched"] == "auto" and feed["discount_names"] == ["DD - Auto - 35% OFF ALL Demo Brand Flower!"]
    assert feed["window"]["discount_total"] == Decimal("715.75") and feed["window"]["transaction_count"] == 33


def test_store_day_totals_verdict_uses_same_weekday_baseline(session):
    from app.models import DailyStoreSummary
    from app.analytics.promotions import promotion_results
    session.add_all([Store(code="HS10001", name="FL - Demo - Pace", state="FL"), Store(code="HS10002", name="FL - Demo - Tampa Kennedy", state="FL")])
    session.commit()
    stores = {s.code: s for s in session.execute(select(Store)).scalars()}
    def day(code, d, rev, gp, disc, tix):
        session.add(DailyStoreSummary(store_id=stores[code].id, sale_date=d, revenue=Decimal(rev), gross_profit=Decimal(gp),
                                      discount_total=Decimal(disc), gross_sales=Decimal(rev) + Decimal(disc), cogs=Decimal(rev) - Decimal(gp), transaction_count=tix))
    # promo Tuesday Sep 15 at both stores; prior Tuesdays Sep 1 and Sep 8 as baseline; a Monday that must be ignored
    day("HS10001", date(2026, 9, 15), "7647.50", "4454.74", "7060.50", 165)
    day("HS10002", date(2026, 9, 15), "9000", "5000", "8000", 200)
    for d in (date(2026, 9, 1), date(2026, 9, 8)):
        day("HS10001", d, "7156.32", "3163.10", "7153.68", 159)
        day("HS10002", d, "8000", "4000", "9000", 180)
    day("HS10001", date(2026, 9, 14), "99999", "99999", "0", 999)
    import_promotion_rows(session, [{"promo": "35% Off Demo Brand Flower", "start_date": "2026-09-15"}])
    pace = promotion_results(session, "HS10001")[0]
    dt = pace["day_totals"]
    assert dt["window"]["revenue_per_store_day"] == Decimal("7647.50") and dt["baseline"]["revenue_per_store_day"] == Decimal("7156.32")
    assert dt["baseline"]["store_days"] == 2 and dt["vs_baseline"]["gross_profit_pct"] == Decimal("0.4083")
    assert dt["verdict"] == "profitable" and pace["verdict"] == "profitable" and "store-day totals" in pace["explanation"]
    fl = promotion_results(session, "state:FL")[0]["day_totals"]
    assert fl["window"]["store_days"] == 2 and fl["window"]["revenue"] == Decimal("16647.50")
    assert fl["baseline"]["store_days"] == 4 and fl["baseline"]["revenue_per_store_day"] == Decimal("7578.16")
    assert promotion_results(session, "HS10002")[0]["day_totals"]["vs_baseline"]["revenue_pct"] == Decimal("0.1250")
